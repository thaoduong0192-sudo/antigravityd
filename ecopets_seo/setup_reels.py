import os
import openpyxl

folder_path = "D:\\Antigravity\\ecopets_seo\\reels_videos"
excel_path = "D:\\Antigravity\\ecopets_seo\\reels_data.xlsx"

# Tạo thư mục nếu chưa có
os.makedirs(folder_path, exist_ok=True)
print(f"Thư mục: {folder_path} đã sẵn sàng.")

# Tạo file Excel nếu chưa có
if not os.path.exists(excel_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reels"
    
    # Tạo Header
    ws.cell(row=1, column=1, value="File Video (.mp4)")
    ws.cell(row=1, column=2, value="Nội dung Caption")
    ws.cell(row=1, column=3, value="Trạng thái")
    
    # Dữ liệu mẫu
    ws.cell(row=2, column=1, value="sample_reel.mp4")
    ws.cell(row=2, column=2, value="Đây là nội dung mẫu cho Reel số 1! Mọi người thấy sao?\n#Ecopets #Reels #Thucung")
    ws.cell(row=2, column=3, value="")
    
    wb.save(excel_path)
    print(f"Đã tạo file Excel mẫu: {excel_path}")
else:
    print(f"File Excel {excel_path} đã tồn tại.")
